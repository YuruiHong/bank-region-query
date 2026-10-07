#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build bank_region_query/data/bank_aliases.json from a personal finance
workbook (default: ~/Documents/balance.xlsx, sheet "银行缩写").

The sheet is laid out as: col A = bank name, cols B..E = 别称 (pinyin
abbreviation, 简称, ...). Each row is one bank, and its cells are grouped, then
that group is assigned to the canonical bank name used by
bank_region_codes.json -- either directly, or by matching "core names" after
dropping a leading 中国 (the sheet writes 中国建设银行 where the data file
writes 建设银行).

Only the sheet's own strings are written out; nothing is derived. The 中国
handling exists purely to decide which data-file key a row belongs to, and the
emitted 别称 are what bank_region_query matches queries against.

Rows that resolve to no known bank are recorded under "unresolved" so the
mapping stays auditable instead of silently dropping data. 别称 claimed by two
different banks land under "conflicts" instead of being written.

Usage:
    python scripts/build_bank_aliases.py [XLSX_PATH] [-o OUTPUT]

Only the standard library plus openpyxl is required.
"""

import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("需要 openpyxl: pip install openpyxl", file=sys.stderr)
    raise

SHEET_NAME = "银行缩写"
DEFAULT_XLSX = Path.home() / "Documents" / "balance.xlsx"
PKG_DATA = Path(__file__).resolve().parent.parent / "bank_region_query" / "data"
DATA_FILE = PKG_DATA / "bank_region_codes.json"
OUTPUT_FILE = PKG_DATA / "bank_aliases.json"

CN_PREFIX = "中国"


def strip_cn_prefix(name: str) -> str:
    """中国建设银行 -> 建设银行. Refuses to shrink a name below 3 chars so
    中国银行 does not collapse to 银行."""
    if name.startswith(CN_PREFIX) and len(name) - len(CN_PREFIX) >= 3:
        return name[len(CN_PREFIX):]
    return name


def load_canonical_names(data_file: Path):
    with open(data_file, "r", encoding="utf-8") as fh:
        return list(json.load(fh).keys())


def read_alias_rows(xlsx_path: Path):
    """Yield (canonical_name, [alias, ...]) per spreadsheet row."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"工作表 '{SHEET_NAME}' 不存在，实际有: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    for row in ws.iter_rows(values_only=True):
        cells = []
        for value in row:
            if value is None:
                continue
            text = str(value).strip()
            if text:
                cells.append(text)
        if cells:
            yield cells[0], cells[1:]


def resolve(candidates, canonical_names):
    """Map a row's name group onto one canonical bank name, or None.

    Two tiers, both scoped to the row's own alias group so 工商银行 can only
    ever resolve to 中国工商银行:
      1. an alias that is literally a data-file key;
      2. a core-name match -- equal after dropping a leading 中国 on both
         sides (中国建设银行 == 建设银行), guarded so 中国银行 cannot decay
         into a bare 银行.
    """
    exact = set(canonical_names)

    for name in candidates:
        if name in exact:
            return name

    cores = {}
    for key in canonical_names:
        core = strip_cn_prefix(key)
        if len(core) >= 3:
            cores.setdefault(core, key)

    for name in candidates:
        core = strip_cn_prefix(name)
        if len(core) >= 3 and core in cores:
            return cores[core]

    return None


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("xlsx", nargs="?", default=str(DEFAULT_XLSX),
                        help=f"balance workbook (default: {DEFAULT_XLSX})")
    parser.add_argument("-o", "--output", default=str(OUTPUT_FILE))
    args = parser.parse_args()

    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        raise SystemExit(f"找不到工作簿: {xlsx_path}")

    canonical_names = load_canonical_names(DATA_FILE)
    aliases = {}
    unresolved = {}

    for canonical_cell, alias_cells in read_alias_rows(xlsx_path):
        candidates = [canonical_cell] + list(alias_cells)
        target = resolve(candidates, canonical_names)
        if target is None:
            unresolved[canonical_cell] = sorted(set(alias_cells) - {canonical_cell})
            continue

        bucket = aliases.setdefault(target, set())
        for name in candidates:
            if name != target:
                bucket.add(name)

    # An alias pointing at two different banks would make lookup ambiguous --
    # drop it and report, rather than pick a winner silently.
    owners = {}
    for target, names in aliases.items():
        for name in names:
            owners.setdefault(name, set()).add(target)

    conflicts = {}
    for name, targets in owners.items():
        if len(targets) > 1:
            conflicts[name] = sorted(targets)
    for target, names in aliases.items():
        aliases[target] = sorted(n for n in names if n not in conflicts)

    payload = {
        "version": 1,
        "generated_from": str(xlsx_path),
        "sheet": SHEET_NAME,
        "canonical_key": "bank_region_codes.json 中的银行名",
        "aliases": {k: aliases[k] for k in sorted(aliases)},
        "unresolved": {k: unresolved[k] for k in sorted(unresolved)},
        "conflicts": {k: conflicts[k] for k in sorted(conflicts)},
    }

    output_path = Path(args.output)
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print(f"写入 {output_path}")
    print(f"  已解析银行 {len(aliases)} / 数据文件共 {len(canonical_names)} 家")
    for target in sorted(aliases):
        if aliases[target]:
            print(f"    {target}: {', '.join(aliases[target])}")
    missing = sorted(set(canonical_names) - set(aliases))
    if missing:
        print(f"  未获得别称的银行 ({len(missing)}): {', '.join(missing)}")
    if unresolved:
        print(f"  未匹配到数据文件的表行 ({len(unresolved)}): {', '.join(unresolved)}")
    if conflicts:
        print(f"  别名冲突已剔除 ({len(conflicts)}): {conflicts}")


if __name__ == "__main__":
    main()
